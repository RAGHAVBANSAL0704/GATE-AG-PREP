import json
import os
import glob
import re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DOCX_DIR = "public/downloads/mock_tests"
os.makedirs(DOCX_DIR, exist_ok=True)

def sanitize_for_docx(s):
    if not isinstance(s, str):
        return str(s) if s is not None else ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', s)

def create_docx_for_mock(mock_num, mock_data):
    m_str = f"{mock_num:02d}"
    docx_path = os.path.join(DOCX_DIR, f"MOCK {m_str} GATE AG.docx")
    doc = docx.Document()

    # Document Header Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = title.add_run(f"GATE 2027\nAgricultural Engineering (AG)\nFull-Length Mock Paper {m_str}\n")
    r1.bold = True
    r1.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Paper Code: AG   |   Duration: 180 minutes (3 hours)\nTotal Questions: 65   |   Maximum Marks: 100\n(Calibrated Strictly to Official GATE AG Recent PYQ Heatmap & Hard Multi-Chain Formulation)\n")
    r2.font.size = Pt(11)
    r2.italic = True

    # General Instructions
    doc.add_heading("GENERAL INSTRUCTIONS", level=2)
    instructions = [
        "1. The paper consists of two sections: General Aptitude (GA, Q.1–Q.10, 15 marks) and Agricultural Engineering (AG, Q.11–Q.65, 85 marks).",
        "2. Q.1–Q.5 and Q.11–Q.35 are 1-mark questions. Q.6–Q.10 and Q.36–Q.65 are 2-mark questions.",
        "3. Questions are of three types: Multiple Choice Questions (MCQ), Multiple Select Questions (MSQ), and Numerical Answer Type (NAT).",
        "4. Marking scheme:\n   • MCQ: +1 or +2 marks for correct answer; -1/3 or -2/3 mark for wrong answer.\n   • MSQ: Full marks only if all correct options are selected and no incorrect option is selected. No negative marking.\n   • NAT: Full marks for numeric answer within acceptable tolerance range. No negative marking.",
        "5. An on-screen virtual calculator is permitted during the examination.",
        "6. All physical quantities follow standard SI units unless explicitly stated otherwise."
    ]
    for inst in instructions:
        p = doc.add_paragraph(inst)
        p.paragraph_format.space_after = Pt(3)

    # 66-Row Summary Table
    doc.add_heading("QUESTION SUMMARY & ANSWER KEY TABLE", level=2)
    table = doc.add_table(rows=66, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Q.No.', 'Section', 'Topic', 'Type', 'Marks', 'Answer']
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True

    questions = mock_data.get('questions', [])
    for idx, q in enumerate(questions):
        row_cells = table.rows[idx + 1].cells
        row_cells[0].text = str(q.get('qnum', idx + 1))
        row_cells[1].text = sanitize_for_docx(q.get('section', ''))
        row_cells[2].text = sanitize_for_docx(q.get('topic', ''))
        row_cells[3].text = str(q.get('type', 'MCQ'))
        row_cells[4].text = str(q.get('marks', 1))
        row_cells[5].text = sanitize_for_docx(str(q.get('correct_answer', '')))

    # Full Question Texts & Detailed Solutions
    doc.add_heading("QUESTIONS & DETAILED STEP-BY-STEP SOLUTIONS", level=2)
    for q in questions:
        qn = q.get('qnum')
        qtype = q.get('type', 'MCQ')
        marks = q.get('marks', 1)
        sec = q.get('section', '')
        top = q.get('topic', '')
        
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(8)
        qr = qp.add_run(f"Q.{qn}  [{sec} | {top}]  ({qtype}, {marks} Mark{'s' if marks>1 else ''})\n")
        qr.bold = True

        doc.add_paragraph(sanitize_for_docx(q.get('question', '')))

        opts = q.get('options')
        if opts and isinstance(opts, dict) and len(opts) > 0:
            for k in sorted(opts.keys()):
                op = doc.add_paragraph(f"({k}) {sanitize_for_docx(str(opts[k]))}")
                op.paragraph_format.left_indent = Inches(0.25)
                op.paragraph_format.space_after = Pt(2)

        ans_p = doc.add_paragraph()
        ans_r = ans_p.add_run(f"Correct Answer: {sanitize_for_docx(str(q.get('correct_answer', '')))}")
        ans_r.bold = True
        ans_r.font.color.rgb = RGBColor(0, 100, 0)

        sol_p = doc.add_paragraph()
        sol_r = sol_p.add_run(f"Solution & Multi-Chain Formula Derivation:\n{sanitize_for_docx(q.get('solution', ''))}\n")
        sol_p.paragraph_format.space_after = Pt(10)

    doc.save(docx_path)
    return docx_path

print("Beginning DOCX sync for all 50 full-length mocks...")
for i in range(1, 51):
    json_path = f"src/data/custom_mock_2027_{i:02d}.json"
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    create_docx_for_mock(i, data)
    print(f"Synced MOCK {i:02d} GATE AG.docx")

print("All 50 DOCX files regenerated and synced successfully!")
