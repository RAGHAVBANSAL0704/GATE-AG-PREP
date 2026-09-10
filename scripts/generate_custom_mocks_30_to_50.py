#!/usr/bin/env python3
"""
Custom Mock Generator: Mocks 30 to 50
Calibrated to 1 notch above actual GATE AG difficulty.
Generates:
1. src/data/custom_mock_2027_XX.json (Mocks 30 to 50)
2. public/downloads/mock_tests/MOCK XX GATE AG.docx (Mocks 30 to 50)
"""

import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import mock_builder as mb

DATA_DIR = "src/data"
DOCX_DIR = "public/downloads/mock_tests"
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(DOCX_DIR, exist_ok=True)

def create_docx_for_mock(mock_num, mock_data):
    m_str = f"{mock_num:02d}"
    docx_path = os.path.join(DOCX_DIR, f"MOCK {m_str} GATE AG.docx")
    doc = docx.Document()

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = title.add_run(f"GATE 2027\nAgricultural Engineering (AG)\nFull-Length Mock Paper {m_str}\n")
    r1.bold = True
    r1.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Paper Code: AG   |   Duration: 180 minutes (3 hours)\nTotal Questions: 65   |   Maximum Marks: 100\n(Calibrated 1 notch above actual GATE AG difficulty with emerging syllabus topics)\n")
    r2.font.size = Pt(11)
    r2.italic = True

    # General Instructions
    doc.add_heading("GENERAL INSTRUCTIONS", level=2)
    instructions = [
        "1. The paper consists of two sections: General Aptitude (GA, Q.1–Q.10, 15 marks) and Agricultural Engineering (AG, Q.11–Q.65, 85 marks).",
        "2. Q.1–Q.5 and Q.11–Q.35 are 1-mark questions. Q.6–Q.10 and Q.36–Q.65 are 2-mark questions.",
        "3. Questions are of three types: Multiple Choice Questions (MCQ), Multiple Select Questions (MSQ), and Numerical Answer Type (NAT).",
        "4. Marking scheme:\n   • MCQ: +1 or +2 marks for correct answer; -1/3 or -2/3 mark for wrong answer.\n   • MSQ: Full marks only if all correct options are selected and no incorrect option is selected. No negative marking.\n   • NAT: Full marks for numeric answer within the specified acceptable range. No negative marking.",
        "5. An on-screen virtual calculator is permitted during the examination.",
        "6. All physical quantities follow standard SI units unless explicitly stated otherwise."
    ]
    for inst in instructions:
        p = doc.add_paragraph(inst)
        p.paragraph_format.space_after = Pt(3)

    # 66-Row Summary Table
    doc.add_heading("QUESTION SUMMARY & ANSWER KEY TABLE", level=2)
    table = doc.add_table(rows=66, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Q.No.', 'Section', 'Sub-topic', 'Type', 'Marks', 'Answer']
    for col_idx, h in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        cell.text = h
        p = cell.paragraphs[0]
        p.runs[0].bold = True

    for q_idx, q in enumerate(mock_data["questions"]):
        row_cells = table.rows[q_idx + 1].cells
        row_cells[0].text = str(q["qnum"])
        row_cells[1].text = q["section"]
        row_cells[2].text = q.get("subtopic", q["topic"])
        row_cells[3].text = q["type"]
        row_cells[4].text = str(q["marks"])
        row_cells[5].text = str(q["correct_answer"])

    doc.add_page_break()

    # Detailed Question Paper with Solutions
    doc.add_heading("QUESTION PAPER & DETAILED SOLUTIONS", level=1)
    for q in mock_data["questions"]:
        sec_header = f"Q.{q['qnum']} ({q['marks']} Mark{'s' if q['marks']>1 else ''}) [{q['type']}] — {q['section']} / {q['topic']}"
        h = doc.add_heading(sec_header, level=3)
        h.paragraph_format.space_before = Pt(8)
        
        qp = doc.add_paragraph(q["question"])
        qp.paragraph_format.space_after = Pt(4)

        if q.get("options") and len(q["options"]) > 0:
            for opt_key in sorted(q["options"].keys()):
                opt_p = doc.add_paragraph(f"({opt_key}) {q['options'][opt_key]}")
                opt_p.paragraph_format.left_indent = Inches(0.25)
                opt_p.paragraph_format.space_after = Pt(2)

        ans_p = doc.add_paragraph()
        r_ans = ans_p.add_run(f"Correct Answer: {q['correct_answer']}\n")
        r_ans.bold = True
        r_sol = ans_p.add_run(f"Detailed Solution:\n{q['solution']}")
        ans_p.paragraph_format.space_after = Pt(8)

    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

def generate_all_mocks():
    for m in range(30, 51):
        m_str = f"{m:02d}"
        json_path = os.path.join(DATA_DIR, f"custom_mock_2027_{m_str}.json")

        ga_qs = mb.get_ga_questions(m)
        t1_qs = mb.get_tech_1m_questions(m)
        t2_qs = mb.get_tech_2m_questions(m)
        all_raw_qs = ga_qs + t1_qs + t2_qs

        processed_qs = []
        for q in all_raw_qs:
            q_num = q["qnum"]
            q_type = q["type"]
            q_marks = q["marks"]
            neg_marks = 0.0
            if q_type == "MCQ":
                neg_marks = 1.0 / 3.0 if q_marks == 1 else 2.0 / 3.0

            q_obj = {
                "id": f"GATE_2027_MOCK_{m_str}_Q{q_num:02d}",
                "year": "2027",
                "qnum": q_num,
                "section": q["section"],
                "topic": q["topic"],
                "subtopic": q.get("subtopic", q["topic"]),
                "gate_section": "GA" if q_num <= 10 else "AG",
                "type": q_type,
                "marks": q_marks,
                "negative_marks": neg_marks,
                "difficulty": q.get("difficulty", "Hard"),
                "question": q["question"],
                "options": q.get("options", {}),
                "correct_answer": q["correct_answer"],
                "solution": q["solution"],
                "isCustomUploaded": True
            }
            processed_qs.append(q_obj)

        mock_data = {
            "id": f"GATE_2027_MOCK_{m_str}",
            "title": f"GATE 2027 AG Full-Length Mock Paper {m_str}",
            "year": "2027",
            "isCustomUploaded": True,
            "instructions": {
                "duration_mins": 180,
                "max_marks": 100,
                "total_qs": 65,
                "instructions": [
                    "Total Duration: 180 minutes (3 hours).",
                    "General Aptitude: 10 Questions (15 Marks).",
                    "Technical Agricultural Engineering: 55 Questions (85 Marks).",
                    "Negative Marking: 1/3 mark deducted for wrong 1-mark MCQ, 2/3 mark for wrong 2-mark MCQ.",
                    "NO negative marking for MSQ or NAT questions."
                ]
            },
            "questions": processed_qs
        }

        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(mock_data, f, indent=2, ensure_ascii=False)
        print(f"Generated JSON: {json_path} (65 Qs, 100 Marks)")

        create_docx_for_mock(m, mock_data)

if __name__ == "__main__":
    generate_all_mocks()
    print("All custom mock tests 30 to 50 successfully generated!")
